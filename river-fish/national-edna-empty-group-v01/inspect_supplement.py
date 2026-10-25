#!/usr/bin/env python3
"""Inspect public Kanagawa eDNA workbooks and emit schema metadata only."""
import hashlib
import io
import json
import re
import urllib.request
import zipfile

SOURCES = {
    "FY2021_water_source_fish": "https://www.pref.kanagawa.jp/documents/97797/r3_kenmin_gyorui.xlsx",
    "FY2022_water_source_fish": "https://www.pref.kanagawa.jp/documents/97797/r4_kenmin_gyorui.xlsx",
    "FY2023_water_source_fish": "https://www.pref.kanagawa.jp/documents/97797/r5_kenmin_gyorui.xlsx",
    "FY2023_prefecture_wide_fish": "https://www.pref.kanagawa.jp/documents/97797/r5_project_gyorui.xlsx",
    "FY2024_water_source_all_taxa": "https://www.pref.kanagawa.jp/documents/97797/r6_kenmin_mtinsects_amphi.xlsx",
    "FY2024_prefecture_wide_all_taxa": "https://www.pref.kanagawa.jp/documents/97797/r6_project_mtinsects-amphi.xlsx",
    "FY2025_water_source_all_taxa": "https://www.pref.kanagawa.jp/documents/97797/r7_kenmin_kekka.xlsx",
    "FY2025_prefecture_wide_all_taxa": "https://www.pref.kanagawa.jp/documents/97797/r7_project_kekka.xlsx",
}
MAX_BYTES = 4_000_000
LABEL = re.compile(
    r"^(県|都道府県|市町村|水系|水系名|河川|河川名|川名|地点名|採水地点|採水地点名|"
    r"調査地点|調査地点名|確認種|確認種名|魚類|魚種|分類群|学名|種名|和名|"
    r"調査日|採水日|調査年度|検出|検出有無|検出結果|出現|出現種|同定結果)$"
)


def inspect(raw, source_id):
    if len(raw) > MAX_BYTES or len(raw) < 1000:
        raise ValueError("Unexpected workbook size")
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    worksheets = []
    for sheet_index, sheet in enumerate(workbook.worksheets, 1):
        header_candidates = []
        max_header_row = min(sheet.max_row or 1, 40)
        for row_number, cells in enumerate(sheet.iter_rows(min_row=1, max_row=max_header_row), 1):
            labels = [(index + 1, str(cell.value).strip())
                      for index, cell in enumerate(cells)
                      if cell.value is not None and LABEL.fullmatch(str(cell.value).strip())]
            if len(labels) >= 2:
                header_candidates.append({"row": row_number, "labels": labels})
        worksheets.append({
            "sheet_index": sheet_index,
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "header_candidates": header_candidates[:8],
        })
    workbook.close()
    return {
        "source_id": source_id,
        "sha256": hashlib.sha256(raw).hexdigest(),
        "bytes": len(raw),
        "worksheets": worksheets,
    }


def fetch_one(source_id, url):
    request = urllib.request.Request(url, headers={
        "User-Agent": "tsurikue-kanagawa-edna-source-inspection/1.0",
        "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "Referer": "https://www.pref.kanagawa.jp/docs/b4f/suigen/edna.html",
    })
    with urllib.request.urlopen(request, timeout=90) as response:
        host = response.geturl().split("/", 3)[2]
        if response.status != 200 or host not in {"www.pref.kanagawa.jp", "www.pref.kanagawa.jp:443"}:
            raise ValueError("Prefecture response left the pinned official host")
        raw = response.read(MAX_BYTES + 1)
    return inspect(raw, source_id)


def fetch_all():
    return {
        "schema": "kanagawa-prefectural-edna-workbook-schema-v01",
        "sources": [fetch_one(source_id, url) for source_id, url in SOURCES.items()],
    }


if __name__ == "__main__":
    print("KANAGAWA_EDNA_SOURCE_SCHEMA", json.dumps(fetch_all(), ensure_ascii=False))
